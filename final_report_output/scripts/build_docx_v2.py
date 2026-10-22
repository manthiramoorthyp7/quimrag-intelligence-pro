"""Builds Final_Project_Report.docx matching the OLD mini-project report's
visual format: Times New Roman, US Letter page size, chapter-opening pages
(CHAPTER N / TITLE centered), ALL-CAPS bold numbered section headings, plain
bordered (non-colored) tables, centered equations, bold centered figure/table
captions, and roman-numeral preliminary-page numbers switching to arabic
numbering at Chapter 1 -- via two Word sections with independent pgNumType.

Usage: python3 build_docx_v2.py report_content.md Final_Project_Report.docx
"""
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLACK = RGBColor(0, 0, 0)
FIG_RE = re.compile(r"\(`([^`]+\.png)`\)")
EQ_RE = re.compile(r"\.\.\. \(\d+\.\d+\)\s*$")


def set_font(run, size=12, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = BLACK
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)


def add_runs(paragraph, text, size=12, base_bold=False):
    tokens = re.split(r"(\*\*.+?\*\*|`[^`]+`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = paragraph.add_run(tok[2:-2])
            set_font(r, size=size, bold=True)
        elif tok.startswith("`") and tok.endswith("`"):
            r = paragraph.add_run(tok[1:-1])
            set_font(r, size=size - 0.5, name="Courier New")
        else:
            r = paragraph.add_run(tok)
            set_font(r, size=size, bold=base_bold)


def set_page_number_field(header_or_footer, fmt="decimal", start=None, align="right"):
    p = header_or_footer.paragraphs[0]
    p.text = ""
    p.alignment = {"right": WD_ALIGN_PARAGRAPH.RIGHT, "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    run = p.add_run()
    set_font(run, size=11)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.append(fld)


def configure_section_numbering(section, fmt, start):
    sectPr = section._sectPr
    pgNumType = sectPr.find(qn("w:pgNumType"))
    if pgNumType is None:
        pgNumType = OxmlElement("w:pgNumType")
        sectPr.append(pgNumType)
    pgNumType.set(qn("w:fmt"), fmt)
    pgNumType.set(qn("w:start"), str(start))


def set_different_first_page(section, enabled=True):
    sectPr = section._sectPr
    titlePg = sectPr.find(qn("w:titlePg"))
    if enabled:
        if titlePg is None:
            titlePg = OxmlElement("w:titlePg")
            sectPr.append(titlePg)
    else:
        if titlePg is not None:
            sectPr.remove(titlePg)
    section.different_first_page_header_footer = enabled


def add_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, h, size=11, base_bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            add_runs(p, val, size=11)
    doc.add_paragraph()


def parse_table_block(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        rows.append(lines[i].strip())
        i += 1
    cells = [[c.strip() for c in r.strip("|").split("|")] for r in rows]
    headers = cells[0]
    body = cells[2:]
    return headers, body, i


def centered(doc, text, size=14, bold=True, space_before=0, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    add_runs(p, text, size=size, base_bold=bold)
    return p


def convert(md_path, docx_path):
    fig_dir = os.path.dirname(os.path.abspath(md_path))
    with open(md_path) as f:
        text = f.read()
    lines = text.split("\n")

    doc = Document()
    style = doc.styles["Normal"]
    set_font_style(style)
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Inches(11), Inches(8.5)
    sec.left_margin = Inches(1.4)
    sec.right_margin = Inches(1.0)
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)
    set_different_first_page(sec, True)  # cover page: no header number
    configure_section_numbering(sec, "lowerRoman", 1)
    set_page_number_field(sec.header, align="right")

    mode = None
    skip_next_break = False
    in_code = False
    code_lines = []
    i = 0
    main_section_started = False

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                in_code = False
                p = doc.add_paragraph()
                box = p.add_run("\n".join(code_lines))
                set_font(box, size=9, name="Courier New")
                pPr = p._p.get_or_add_pPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:fill"), "F2F2F2")
                pPr.append(shd)
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if stripped == "<!--COVER-->":
            mode = "cover"
            i += 1
            continue
        if stripped == "<!--CERTIFICATE-->":
            mode = "certificate"
            i += 1
            continue
        if stripped == "<!--ACK-->":
            mode = "ack"
            i += 1
            continue
        if stripped == "<!--ABSTRACT-->":
            mode = "abstract"
            i += 1
            continue
        if stripped == "<!--MAINSTART-->":
            new_sec = doc.add_section(WD_SECTION.NEW_PAGE)
            new_sec.left_margin = Inches(1.4)
            new_sec.right_margin = Inches(1.0)
            new_sec.top_margin = Inches(1.0)
            new_sec.bottom_margin = Inches(1.0)
            set_different_first_page(new_sec, False)
            configure_section_numbering(new_sec, "decimal", 1)
            set_page_number_field(new_sec.header, align="right")
            main_section_started = True
            mode = "body"
            skip_next_break = True
            i += 1
            continue
        if stripped.startswith("<!--CHAPTER:"):
            m = re.match(r"<!--CHAPTER:([^:]*):(.*)-->", stripped)
            chno, chtitle = m.group(1), m.group(2)
            if chno.strip() and not skip_next_break:
                doc.add_page_break()
            skip_next_break = False
            if chno == "REFERENCES":
                centered(doc, "REFERENCES", size=16, bold=True, space_before=0, space_after=18)
            else:
                centered(doc, f"CHAPTER {chno}", size=16, bold=True, space_before=0, space_after=8)
                centered(doc, chtitle, size=16, bold=True, space_before=0, space_after=18)
            mode = "body"
            i += 1
            continue
        if stripped == "<!--PAGEBREAK-->":
            doc.add_page_break()
            i += 1
            continue

        if not stripped:
            i += 1
            continue
        if stripped == "---":
            i += 1
            continue

        if mode == "cover" and stripped and not stripped.startswith("("):
            # First substantial line is the long title; subsequent lines are
            # short centered metadata lines.
            is_title = len(stripped) > 60
            centered(doc, stripped, size=18 if is_title else 13,
                     bold=is_title, space_before=6 if not is_title else 40,
                     space_after=10 if is_title else 8)
            i += 1
            continue
        if mode == "cover" and stripped.startswith("("):
            centered(doc, stripped, size=12, bold=False, space_after=8)
            i += 1
            continue

        if mode in ("certificate", "ack", "abstract") and stripped in (
                "CERTIFICATE", "ACKNOWLEDGEMENT", "ABSTRACT", "TABLE OF CONTENTS",
                "LIST OF TABLES", "LIST OF FIGURES", "LIST OF ABBREVIATIONS"):
            centered(doc, stripped, size=15, bold=True, space_after=16)
            i += 1
            continue

        if stripped in ("TABLE OF CONTENTS", "LIST OF TABLES", "LIST OF FIGURES", "LIST OF ABBREVIATIONS"):
            centered(doc, stripped, size=15, bold=True, space_after=16)
            i += 1
            continue

        if stripped.startswith("### "):
            txt = stripped[4:]
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.left_indent = Inches(0.15)
            r = h.add_run(txt.rstrip(":") + ":" if not txt.endswith(":") else txt)
            set_font(r, size=12, bold=True)
            i += 1
            continue
        if stripped.startswith("## "):
            txt = stripped[3:]
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            r = h.add_run(txt)
            set_font(r, size=13, bold=True)
            i += 1
            continue

        if stripped.startswith("| "):
            headers, body, nexti = parse_table_block(lines, i)
            add_table(doc, headers, body)
            i = nexti
            continue

        if re.match(r"^[-*] ", stripped):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, stripped[2:], size=12)
            i += 1
            continue
        if re.match(r"^\d+\. ", stripped):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+\. ", "", stripped), size=12)
            i += 1
            continue

        fig_match = FIG_RE.search(stripped)
        if fig_match and stripped.startswith("**Fig"):
            img_rel = fig_match.group(1)
            img_path = os.path.join(fig_dir, img_rel)
            caption_text = FIG_RE.sub("", stripped).strip()
            if os.path.exists(img_path):
                doc.add_picture(img_path, width=Inches(5.6))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            centered(doc, caption_text, size=11, bold=True, space_after=10)
            i += 1
            continue

        if EQ_RE.search(stripped):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            add_runs(p, stripped, size=12)
            i += 1
            continue

        if stripped.startswith("Where:") or stripped.startswith("Where "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_runs(p, stripped, size=12)
            i += 1
            continue

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.15
        add_runs(p, stripped, size=12)
        i += 1

    doc.save(docx_path)
    print(f"wrote {docx_path}")


def set_font_style(style):
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.font.color.rgb = BLACK
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "Times New Roman")


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
